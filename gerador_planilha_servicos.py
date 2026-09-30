"""
Módulo: gerador_planilha_servicos.py
Gera o Excel detalhado do certame de serviço aprovado.
"""
from pathlib import Path
from decimal import Decimal
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PASTA_DESTINO = Path.home() / "Documents" / "Licitacoes_Aprovadas"
PASTA_DESTINO.mkdir(parents=True, exist_ok=True)

def exportar_pregao_servico_excel(dados_licitacao: dict, analise_operacional: dict):
    """
    Gera o Excel executivo com custos, logística, equipe e ferramentas.
    """
    wb = Workbook()
    
    # -------------------------------------------------------------
    # ABA 1: FORMAÇÃO DE PREÇO & DADOS DO EDITAL
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Resumo e Preços"
    
    cabecalho_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font_branca = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    ws1.append(["CAMPO", "DETALHAMENTO"])
    ws1.append(["Órgão Licitante", dados_licitacao.get("orgao_nome")])
    ws1.append(["Edital / Processo", dados_licitacao.get("numero_edital")])
    ws1.append(["Cidade / UF de Execução", f"{dados_licitacao.get('municipio')} - {dados_licitacao.get('uf')}"])
    ws1.append(["Link Direto PNCP", dados_licitacao.get("link_edital")])
    ws1.append(["", ""])
    ws1.append(["Valor Referência (Governo)", f"R$ {dados_licitacao.get('valor_referencia', 0):,.2f}"])
    ws1.append(["Custo Operacional Total", f"R$ {analise_operacional.get('custo_total', 0):,.2f}"])
    ws1.append(["Piso Mínimo JRBC (Margem >= 30%)", f"R$ {analise_operacional.get('piso_minimo', 0):,.2f}"])
    ws1.append(["Margem de Lucro Estimada", f"{analise_operacional.get('margem_pct', 0):.2f}%"])
    ws1.append(["Lucro Líquido Projetado", f"R$ {analise_operacional.get('lucro_projetado', 0):,.2f}"])
    
    # -------------------------------------------------------------
    # ABA 2: LOGÍSTICA & DISTÂNCIA (BASE: PALMAS/TO)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Logística e Deslocamento")
    ws2.append(["Parâmetro", "Valor Estimado"])
    ws2.append(["Origem da Operação", "Palmas - TO"])
    ws2.append(["Destino da Execução", f"{dados_licitacao.get('municipio')} - {dados_licitacao.get('uf')}"])
    ws2.append(["Distância Estimada (Ida e Volta)", f"{analise_operacional.get('distancia_km', 0)} km"])
    ws2.append(["Custo com Combustível / Passagem", f"R$ {analise_operacional.get('custo_transporte', 0):,.2f}"])
    ws2.append(["Diárias / Alimentação / Hospedagem", f"R$ {analise_operacional.get('custo_diarias', 0):,.2f}"])
    ws2.append(["Modalidade de Atendimento", analise_operacional.get("tipo_atendimento", "Presencial")])
    
    # -------------------------------------------------------------
    # ABA 3: EQUIPE E FERRAMENTAL
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Equipe e Ferramentas")
    ws3.append(["Equipe Necessária", "Horas Estimadas / Mês", "Função"])
    for membro in analise_operacional.get("equipe", []):
        ws3.append([membro["nome"], membro["horas"], membro["funcao"]])
        
    ws3.append(["", "", ""])
    ws3.append(["Ferramental / Equipamentos Requeridos", "Finalidade", "Disponibilidade"])
    for ferr in analise_operacional.get("ferramentas", []):
        ws3.append([ferr["item"], ferr["finalidade"], ferr["status"]])

    # Ajuste de largura das colunas
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    nome_arquivo = f"Servico_{dados_licitacao.get('uf')}_{dados_licitacao.get('municipio')}_{dados_licitacao.get('numero_edital', '0').replace('/', '-')}.xlsx"
    caminho_final = PASTA_DESTINO / nome_arquivo
    wb.save(caminho_final)
    return str(caminho_final)
